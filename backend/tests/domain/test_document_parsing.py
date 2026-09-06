from __future__ import annotations

import json
import zipfile
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from ych_bot.application import DocumentParseError, DocumentTextExtractor


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("User notes", level=1)
    document.add_paragraph("Likes street photography and quiet cafés.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Communication"
    table.cell(0, 1).text = "concise"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def text_pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def test_docx_extracts_paragraphs_and_tables() -> None:
    parsed = DocumentTextExtractor().extract(filename="notes.docx", content=docx_bytes())

    assert parsed.detected_format == "docx"
    assert "Likes street photography" in parsed.text
    assert "Communication\tconcise" in parsed.text
    assert parsed.metadata["table_count"] == 1


class FakePdfOcrEngine:
    def __init__(self, pages: dict[int, str] | None = None) -> None:
        self.calls: list[int] = []
        self.pages = pages or {}

    def recognize_page(self, *, page_index: int, pdf_bytes: bytes) -> str:
        del pdf_bytes
        self.calls.append(page_index)
        return self.pages.get(page_index, "")


def blank_pdf_bytes(pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_pdf_extracts_page_markers_and_text() -> None:
    parsed = DocumentTextExtractor().extract(
        filename="diary.pdf",
        content=text_pdf_bytes("A calm day and a long walk"),
    )

    assert parsed.detected_format == "pdf_text"
    assert "[PDF page 1]" in parsed.text
    assert "A calm day and a long walk" in parsed.text
    assert parsed.metadata == {
        "page_count": 1,
        "text_page_count": 1,
        "ocr_page_count": 0,
        "ocr_enabled": False,
    }


def test_scanned_pdf_does_not_call_ocr_when_disabled() -> None:
    engine = FakePdfOcrEngine({1: "should not be used"})
    extractor = DocumentTextExtractor(ocr_enabled=False, ocr_engine=engine)

    with pytest.raises(DocumentParseError, match="OCR is disabled"):
        extractor.extract(filename="scan.pdf", content=blank_pdf_bytes())
    assert engine.calls == []


def test_scanned_pdf_uses_ocr_only_for_pages_without_text() -> None:
    engine = FakePdfOcrEngine({1: "handwritten cafe notes"})
    extractor = DocumentTextExtractor(ocr_enabled=True, ocr_engine=engine)
    scanned = extractor.extract(filename="scan.pdf", content=blank_pdf_bytes())
    text_pdf = extractor.extract(
        filename="diary.pdf",
        content=text_pdf_bytes("A calm day and a long walk"),
    )

    assert scanned.detected_format == "pdf_ocr"
    assert scanned.text == "[PDF page 1][OCR]\nhandwritten cafe notes"
    assert scanned.metadata == {
        "page_count": 1,
        "text_page_count": 0,
        "ocr_page_count": 1,
        "ocr_enabled": True,
    }
    assert engine.calls == [1]
    assert text_pdf.detected_format == "pdf_text"
    assert "[OCR]" not in text_pdf.text
    assert engine.calls == [1]


def test_mixed_pdf_ocrs_only_empty_text_layer_pages() -> None:
    writer = PdfWriter()
    writer.add_page(PdfReader(BytesIO(text_pdf_bytes("A calm day"))).pages[0])
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    engine = FakePdfOcrEngine({2: "margin doodles"})
    parsed = DocumentTextExtractor(ocr_enabled=True, ocr_engine=engine).extract(
        filename="mixed.pdf",
        content=buffer.getvalue(),
    )

    assert parsed.detected_format == "pdf_ocr"
    assert "[PDF page 1]\nA calm day" in parsed.text
    assert "[PDF page 2][OCR]\nmargin doodles" in parsed.text
    assert parsed.metadata == {
        "page_count": 2,
        "text_page_count": 1,
        "ocr_page_count": 1,
        "ocr_enabled": True,
    }
    assert engine.calls == [2]


def qq_pc_txt_export() -> bytes:
    return "\n".join(
        [
            "消息记录（此消息记录为文本格式，不支持重新导入）",
            "",
            "================================================",
            "消息分组:我的好友",
            "================================================",
            "消息对象:Alice",
            "================================================",
            "",
            "2026-08-13 10:00:00 Alice(10001)",
            "hello cafe",
            "see you later",
            "",
            "2026/8/13 10:01:00 我",
            "在的",
        ]
    ).encode("utf-8")


def test_qq_txt_export_is_normalized_without_forcing_plain_notes() -> None:
    extractor = DocumentTextExtractor()
    chat = extractor.extract(filename="alice.txt", content=qq_pc_txt_export())
    note = extractor.extract(
        filename="notes.txt",
        content=b"Meeting notes\nBuy coffee after 2026-08-13 10:00:00 if free.\n",
    )

    assert chat.detected_format == "chat_txt"
    assert chat.text == (
        "[2026-08-13 10:00:00] Alice(10001): hello cafe\nsee you later\n"
        "[2026/8/13 10:01:00] 我: 在的"
    )
    assert chat.metadata == {"message_count": 2, "streamed": False}
    assert note.detected_format == "text"
    assert "Buy coffee" in note.text
    assert note.metadata == {}


def test_qq_txt_export_respects_record_limit() -> None:
    with pytest.raises(DocumentParseError, match="record limit"):
        DocumentTextExtractor(max_records=1).extract(
            filename="alice.txt",
            content=qq_pc_txt_export(),
        )


def test_onebot_json_uses_raw_message_nested_sender_and_unix_time() -> None:
    timestamp = int(datetime(2026, 8, 13, 10, 0, 0, tzinfo=ZoneInfo("Asia/Shanghai")).timestamp())
    payload = json.dumps(
        [
            {
                "time": timestamp,
                "sender": {"user_id": 10001, "nickname": "Alice"},
                "raw_message": "hello cafe",
                "message": [{"type": "text", "data": {"text": "hello cafe"}}],
            }
        ]
    )
    parsed = DocumentTextExtractor().extract(filename="onebot.json", content=payload.encode())

    assert parsed.detected_format == "chat_json"
    assert parsed.text == "[2026-08-13 10:00:00] Alice: hello cafe"
    assert parsed.metadata == {"message_count": 1, "streamed": True}


def test_qq_html_export_is_normalized_while_plain_html_stays_visible_text() -> None:
    extractor = DocumentTextExtractor()
    chat = extractor.extract(
        filename="export.html",
        content="".join(
            [
                "<html><body>",
                '<div class="message">',
                '<span class="time">2026-08-13 10:00:00</span>',
                '<span class="nickname">Alice</span>',
                '<div class="content">hello cafe</div>',
                "</div>",
                '<div class="message">',
                '<span class="msg-time">2026-08-13 10:01:00</span>',
                '<span class="msg-nick">Bob</span>',
                '<div class="msg-content">在的</div>',
                "</div>",
                "</body></html>",
            ]
        ).encode(),
    )
    page = extractor.extract(
        filename="notes.html",
        content=b"<style>hidden</style><p>Alice: visible</p><script>bad()</script>",
    )

    assert chat.detected_format == "chat_html"
    assert chat.text == "[2026-08-13 10:00:00] Alice: hello cafe\n[2026-08-13 10:01:00] Bob: 在的"
    assert chat.metadata == {"message_count": 2, "streamed": False}
    assert page.detected_format == "html"
    assert page.text == "Alice: visible"


def test_chat_json_csv_and_html_are_normalized() -> None:
    extractor = DocumentTextExtractor()
    json_result = extractor.extract(
        filename="chat.json",
        content=b'[{"timestamp":"2026-08-13 10:00","sender":"Alice","message":"hello"}]',
    )
    csv_result = extractor.extract(
        filename="chat.csv",
        content="时间,昵称,内容\n10:01,Bob,你好".encode(),
    )
    html_result = extractor.extract(
        filename="export.html",
        content=b"<style>hidden</style><p>Alice: visible</p><script>bad()</script>",
    )

    assert json_result.detected_format == "chat_json"
    assert json_result.text == "[2026-08-13 10:00] Alice: hello"
    assert csv_result.detected_format == "chat_csv"
    assert csv_result.text == "[10:01] Bob: 你好"
    assert html_result.detected_format == "html"
    assert html_result.text == "Alice: visible"
    assert "hidden" not in html_result.text
    assert "bad" not in html_result.text
    assert json_result.metadata == {"message_count": 1, "streamed": True}
    assert csv_result.metadata == {"message_count": 1, "streamed": True}


def test_wrapped_chat_json_is_streamed_without_loading_all_rows() -> None:
    messages = [
        {"timestamp": f"2026-08-13 10:{index:02d}", "sender": "Alice", "message": f"line {index}"}
        for index in range(20)
    ]
    payload = json.dumps({"meta": {"source": "qq"}, "messages": messages}, ensure_ascii=False)
    parsed = DocumentTextExtractor(max_records=50).extract(
        filename="export.json",
        content=payload.encode(),
    )

    assert parsed.detected_format == "chat_json"
    assert parsed.metadata == {"message_count": 20, "streamed": True}
    assert parsed.text.startswith("[2026-08-13 10:00] Alice: line 0")
    assert parsed.text.endswith("[2026-08-13 10:19] Alice: line 19")


def test_chat_exports_respect_record_and_eager_json_limits() -> None:
    extractor = DocumentTextExtractor(max_records=2, max_eager_json_bytes=64)
    oversized_chat = json.dumps(
        [
            {"timestamp": "10:00", "sender": "A", "message": "one"},
            {"timestamp": "10:01", "sender": "B", "message": "two"},
            {"timestamp": "10:02", "sender": "C", "message": "three"},
        ]
    )
    oversized_csv = "时间,昵称,内容\n10:00,A,one\n10:01,B,two\n10:02,C,three\n"
    bulky_object = json.dumps({"notes": "x" * 80})

    with pytest.raises(DocumentParseError, match="record limit"):
        extractor.extract(filename="chat.json", content=oversized_chat.encode())
    with pytest.raises(DocumentParseError, match="record limit"):
        extractor.extract(filename="chat.csv", content=oversized_csv.encode())
    with pytest.raises(DocumentParseError, match="too large to parse eagerly"):
        extractor.extract(filename="notes.json", content=bulky_object.encode())


def test_docx_zip_bomb_and_pdf_page_limit_are_rejected() -> None:
    archive_buffer = BytesIO()
    with zipfile.ZipFile(archive_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * 2_000_000)
    with pytest.raises(DocumentParseError, match="compression ratio"):
        DocumentTextExtractor().extract(
            filename="bomb.docx",
            content=archive_buffer.getvalue(),
        )

    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_blank_page(width=100, height=100)
    pdf_buffer = BytesIO()
    writer.write(pdf_buffer)
    with pytest.raises(DocumentParseError, match="page limit"):
        DocumentTextExtractor(max_pages=1).extract(
            filename="too-many.pdf",
            content=pdf_buffer.getvalue(),
        )
