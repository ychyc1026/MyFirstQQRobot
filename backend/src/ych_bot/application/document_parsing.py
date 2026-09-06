"""Bounded text extraction for long-form knowledge imports."""

from __future__ import annotations

import csv
import json
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import datetime
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Protocol
from zoneinfo import ZoneInfo

from docx import Document
from pypdf import PdfReader


class DocumentParseError(ValueError):
    pass


class PdfOcrEngine(Protocol):
    def recognize_page(self, *, page_index: int, pdf_bytes: bytes) -> str: ...


class DisabledPdfOcrEngine:
    def recognize_page(self, *, page_index: int, pdf_bytes: bytes) -> str:
        del page_index, pdf_bytes
        raise DocumentParseError("PDF OCR engine is not configured")


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    text: str
    detected_format: str
    metadata: dict[str, object]


class DocumentTextExtractor:
    supported_suffixes = {
        ".txt",
        ".md",
        ".json",
        ".csv",
        ".html",
        ".htm",
        ".docx",
        ".pdf",
    }

    def __init__(
        self,
        *,
        max_pages: int = 2000,
        max_archive_entries: int = 2000,
        max_expanded_bytes: int = 500 * 1024 * 1024,
        max_records: int = 100_000,
        max_eager_json_bytes: int = 5 * 1024 * 1024,
        ocr_enabled: bool = False,
        ocr_engine: PdfOcrEngine | None = None,
    ) -> None:
        if max_pages < 1:
            raise ValueError("document max pages must be positive")
        if max_archive_entries < 1 or max_expanded_bytes < 1:
            raise ValueError("document archive limits must be positive")
        if max_records < 1:
            raise ValueError("document max records must be positive")
        if max_eager_json_bytes < 1:
            raise ValueError("document eager JSON limit must be positive")
        self._max_pages = max_pages
        self._max_archive_entries = max_archive_entries
        self._max_expanded_bytes = max_expanded_bytes
        self._max_records = max_records
        self._max_eager_json_bytes = max_eager_json_bytes
        self._ocr_enabled = ocr_enabled
        self._ocr_engine = ocr_engine

    def extract(self, *, filename: str, content: bytes) -> ParsedDocument:
        suffix = Path(filename).suffix.lower()
        if suffix not in self.supported_suffixes:
            supported = ", ".join(sorted(self.supported_suffixes))
            raise DocumentParseError(f"unsupported document type; supported: {supported}")
        if suffix == ".docx":
            return self._extract_docx(content)
        if suffix == ".pdf":
            return self._extract_pdf(content)
        text = _decode_text(content)
        if suffix in {".html", ".htm"}:
            chat = _normalize_qq_html(
                text,
                max_records=self._max_records,
                max_expanded_bytes=self._max_expanded_bytes,
            )
            if chat is not None:
                return self._checked(
                    ParsedDocument(
                        text=chat[0],
                        detected_format="chat_html",
                        metadata=_chat_metadata(chat[1], False),
                    )
                )
            return self._checked(
                ParsedDocument(
                    text=_extract_html(text),
                    detected_format="html",
                    metadata={},
                )
            )
        if suffix == ".json":
            normalized, count, streamed = _normalize_json(
                text,
                max_records=self._max_records,
                max_expanded_bytes=self._max_expanded_bytes,
                max_eager_json_bytes=self._max_eager_json_bytes,
            )
            return self._checked(
                ParsedDocument(
                    text=normalized,
                    detected_format="chat_json" if count is not None else "json",
                    metadata=_chat_metadata(count, streamed),
                )
            )
        if suffix == ".csv":
            normalized, count, streamed = _normalize_csv(
                text,
                max_records=self._max_records,
                max_expanded_bytes=self._max_expanded_bytes,
            )
            return self._checked(
                ParsedDocument(
                    text=normalized,
                    detected_format="chat_csv" if count is not None else "csv",
                    metadata=_chat_metadata(count, streamed),
                )
            )
        if suffix == ".txt":
            chat = _normalize_qq_txt(
                text,
                max_records=self._max_records,
                max_expanded_bytes=self._max_expanded_bytes,
            )
            if chat is not None:
                return self._checked(
                    ParsedDocument(
                        text=chat[0],
                        detected_format="chat_txt",
                        metadata=_chat_metadata(chat[1], False),
                    )
                )
        return self._checked(
            ParsedDocument(
                text=text,
                detected_format="markdown" if suffix == ".md" else "text",
                metadata={},
            )
        )

    def _extract_docx(self, content: bytes) -> ParsedDocument:
        self._validate_docx_archive(content)
        try:
            document = Document(BytesIO(content))
        except Exception as exc:
            raise DocumentParseError("DOCX could not be parsed") from exc
        blocks: list[str] = []
        blocks.extend(paragraph.text.strip() for paragraph in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    blocks.append("\t".join(values))
        for section in document.sections:
            blocks.extend(paragraph.text.strip() for paragraph in section.header.paragraphs)
            blocks.extend(paragraph.text.strip() for paragraph in section.footer.paragraphs)
        text = "\n\n".join(value for value in blocks if value)
        self._validate_expanded_text(text)
        return ParsedDocument(
            text=text,
            detected_format="docx",
            metadata={
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            },
        )

    def _validate_docx_archive(self, content: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
        except (zipfile.BadZipFile, OSError) as exc:
            raise DocumentParseError("DOCX is not a valid ZIP package") from exc
        if len(entries) > self._max_archive_entries:
            raise DocumentParseError("DOCX contains too many archive entries")
        expanded = 0
        for entry in entries:
            if entry.flag_bits & 0x1:
                raise DocumentParseError("encrypted DOCX archives are not supported")
            expanded += entry.file_size
            if expanded > self._max_expanded_bytes:
                raise DocumentParseError("DOCX expanded content exceeds the configured limit")
            if entry.compress_size and entry.file_size > max(entry.compress_size * 200, 1_000_000):
                raise DocumentParseError("DOCX contains a suspicious compression ratio")

    def _extract_pdf(self, content: bytes) -> ParsedDocument:
        try:
            reader = PdfReader(BytesIO(content), strict=False)
        except Exception as exc:
            raise DocumentParseError("PDF could not be parsed") from exc
        if reader.is_encrypted:
            raise DocumentParseError("encrypted PDF files are not supported")
        if len(reader.pages) > self._max_pages:
            raise DocumentParseError("PDF exceeds the configured page limit")
        pages: list[str] = []
        expanded_bytes = 0
        text_page_count = 0
        ocr_page_count = 0
        for index, page in enumerate(reader.pages, start=1):
            try:
                extracted = page.extract_text() or ""
            except Exception as exc:
                raise DocumentParseError(f"PDF page {index} could not be extracted") from exc
            cleaned = extracted.replace("\x00", "").strip()
            source = "text"
            if not cleaned and self._ocr_enabled:
                engine = self._ocr_engine or DisabledPdfOcrEngine()
                cleaned = (
                    engine.recognize_page(page_index=index, pdf_bytes=content)
                    .replace("\x00", "")
                    .strip()
                )
                source = "ocr"
            if not cleaned:
                continue
            marker = f"[PDF page {index}][OCR]" if source == "ocr" else f"[PDF page {index}]"
            block = f"{marker}\n{cleaned}"
            expanded_bytes += len(block.encode("utf-8"))
            if expanded_bytes > self._max_expanded_bytes:
                raise DocumentParseError("extracted document text exceeds the configured limit")
            pages.append(block)
            if source == "ocr":
                ocr_page_count += 1
            else:
                text_page_count += 1
        if not pages:
            if self._ocr_enabled:
                raise DocumentParseError("OCR produced no readable text")
            raise DocumentParseError("scanned PDF has no readable text layer; OCR is disabled")
        text = "\n\n".join(pages)
        self._validate_expanded_text(text)
        return ParsedDocument(
            text=text,
            detected_format="pdf_ocr" if ocr_page_count else "pdf_text",
            metadata={
                "page_count": len(reader.pages),
                "text_page_count": text_page_count,
                "ocr_page_count": ocr_page_count,
                "ocr_enabled": self._ocr_enabled,
            },
        )

    def _validate_expanded_text(self, text: str) -> None:
        if len(text.encode("utf-8")) > self._max_expanded_bytes:
            raise DocumentParseError("extracted document text exceeds the configured limit")

    def _checked(self, parsed: ParsedDocument) -> ParsedDocument:
        self._validate_expanded_text(parsed.text)
        return parsed


def _decode_text(content: bytes) -> str:
    if b"\x00" in content:
        raise DocumentParseError("binary content is not accepted for text import")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError("text document must use UTF-8 encoding") from exc


class _VisibleHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._ignored_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag.lower() in {"script", "style", "noscript"}:
            self._ignored_depth += 1
        elif tag.lower() in {"br", "p", "div", "li", "tr"} and not self._ignored_depth:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag.lower() in {"p", "div", "li", "tr"} and not self._ignored_depth:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth:
            self.parts.append(data)


def _extract_html(text: str) -> str:
    parser = _VisibleHTMLParser()
    try:
        parser.feed(text)
        parser.close()
    except Exception as exc:
        raise DocumentParseError("HTML could not be parsed") from exc
    return "\n".join(line.strip() for line in "".join(parser.parts).splitlines() if line.strip())


def _normalize_qq_txt(
    text: str,
    *,
    max_records: int,
    max_expanded_bytes: int,
) -> tuple[str, int] | None:
    has_header = any(marker in text for marker in _QQ_TXT_HEADER_MARKERS)
    messages: list[tuple[str, str, list[str]]] = []
    current: tuple[str, str, list[str]] | None = None
    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        stripped = raw_line.strip()
        if not stripped or set(stripped) <= {"=", "-", "*"}:
            continue
        if any(stripped.startswith(prefix) for prefix in _QQ_TXT_HEADER_MARKERS):
            continue
        match = _QQ_TXT_START.match(stripped)
        if match is not None:
            if current is not None:
                messages.append(current)
            current = (match.group(1), match.group(2).strip(), [])
            continue
        if current is not None:
            current[2].append(stripped)
    if current is not None:
        messages.append(current)
    usable = [(timestamp, sender, body) for timestamp, sender, body in messages if body]
    if not usable or (not has_header and len(usable) < 2):
        return None
    return _join_chat_lines(
        [f"[{timestamp}] {sender}: " + "\n".join(body) for timestamp, sender, body in usable],
        max_records=max_records,
        max_expanded_bytes=max_expanded_bytes,
        limit_label="TXT chat export",
    )


def _normalize_qq_html(
    text: str,
    *,
    max_records: int,
    max_expanded_bytes: int,
) -> tuple[str, int] | None:
    parser = _QQHtmlChatParser()
    try:
        parser.feed(text)
        parser.close()
    except Exception as exc:
        raise DocumentParseError("HTML could not be parsed") from exc
    usable = [
        item for item in parser.messages if item["content"] and (item["time"] or item["sender"])
    ]
    if not usable:
        return None
    return _join_chat_lines(
        [_format_message_row(item, "time", "sender", "content") for item in usable],
        max_records=max_records,
        max_expanded_bytes=max_expanded_bytes,
        limit_label="HTML chat export",
    )


def _join_chat_lines(
    lines: list[str],
    *,
    max_records: int,
    max_expanded_bytes: int,
    limit_label: str,
) -> tuple[str, int]:
    if len(lines) > max_records:
        raise DocumentParseError(f"{limit_label} exceeds the configured record limit")
    expanded = 0
    output: list[str] = []
    for line in lines:
        expanded += len(line.encode("utf-8")) + (1 if output else 0)
        if expanded > max_expanded_bytes:
            raise DocumentParseError("extracted document text exceeds the configured limit")
        output.append(line)
    return "\n".join(output), len(output)


def _class_tokens(attrs: list[tuple[str, str | None]]) -> set[str]:
    for key, value in attrs:
        if key.lower() == "class" and value:
            return {token.lower().replace("_", "-") for token in value.split() if token}
    return set()


class _QQHtmlChatParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.messages: list[dict[str, str]] = []
        self._depth = 0
        self._field: str | None = None
        self._field_depth = 0
        self._current: dict[str, str] | None = None
        self._buffer: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._ignored_depth += 1
            return
        if self._ignored_depth:
            return
        classes = _class_tokens(attrs)
        if self._current is None:
            if classes & _MESSAGE_CONTAINER_CLASSES:
                self._current = {"time": "", "sender": "", "content": ""}
                self._depth = 1
            return
        self._depth += 1
        if self._field is not None:
            return
        if classes & _TIME_CLASSES:
            field = "time"
        elif classes & _SENDER_CLASSES:
            field = "sender"
        elif classes & _CONTENT_CLASSES:
            field = "content"
        else:
            return
        self._field = field
        self._field_depth = self._depth
        self._buffer = []

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
            return
        if self._ignored_depth or self._current is None:
            return
        if self._field is not None and self._depth == self._field_depth:
            self._flush_field()
            self._field = None
        self._depth -= 1
        if self._depth <= 0:
            self._flush_field()
            if self._current["content"] or self._current["sender"] or self._current["time"]:
                self.messages.append(self._current)
            self._current = None
            self._field = None
            self._depth = 0

    def handle_data(self, data: str) -> None:
        if self._ignored_depth or self._field is None:
            return
        self._buffer.append(data)

    def _flush_field(self) -> None:
        if self._current is None or self._field is None:
            return
        text = " ".join(part.strip() for part in "".join(self._buffer).split() if part.strip())
        if text:
            existing = self._current[self._field]
            self._current[self._field] = f"{existing} {text}".strip() if existing else text
        self._buffer = []


_SHANGHAI = ZoneInfo("Asia/Shanghai")
_QQ_TXT_HEADER_MARKERS = ("消息记录", "消息分组", "消息对象")
_QQ_TXT_START = re.compile(r"^(\d{4}[-/]\d{1,2}[-/]\d{1,2}\s+\d{1,2}:\d{2}(?::\d{2})?)\s+(\S.*)$")
_MESSAGE_CONTAINER_CLASSES = {
    "message",
    "msg",
    "chat-item",
    "msg-item",
    "message-item",
}
_TIME_CLASSES = {"time", "msg-time", "message-time", "chat-time", "timestamp", "msgtime"}
_SENDER_CLASSES = {
    "nickname",
    "nick",
    "msg-nick",
    "sender",
    "name",
    "username",
    "msg-nickname",
}
_CONTENT_CLASSES = {
    "content",
    "msg-content",
    "message-content",
    "text",
    "body",
    "msg-text",
}
_TIME_KEYS = (
    "time",
    "timestamp",
    "datetime",
    "date",
    "send_time",
    "msg_time",
    "msgtime",
    "create_time",
    "ts",
    "时间",
)
_SENDER_KEYS = (
    "sender_name",
    "nickname",
    "nick",
    "name",
    "sender",
    "user",
    "username",
    "from",
    "qq",
    "uin",
    "user_id",
    "from_id",
    "发送者",
    "昵称",
)
_CONTENT_KEYS = (
    "raw_message",
    "msg_content",
    "content",
    "text",
    "body",
    "message",
    "msg",
    "消息",
    "内容",
)


def _chat_metadata(count: int | None, streamed: bool) -> dict[str, object]:
    if count is None:
        return {}
    return {"message_count": count, "streamed": streamed}


def _normalize_json(
    text: str,
    *,
    max_records: int,
    max_expanded_bytes: int,
    max_eager_json_bytes: int,
) -> tuple[str, int | None, bool]:
    streamed = _stream_chat_json(
        text,
        max_records=max_records,
        max_expanded_bytes=max_expanded_bytes,
    )
    if streamed is not None:
        return streamed
    if len(text.encode("utf-8")) > max_eager_json_bytes:
        raise DocumentParseError("JSON document is too large to parse eagerly")
    try:
        value = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DocumentParseError("JSON document is invalid") from exc
    return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True), None, False


def _normalize_csv(
    text: str,
    *,
    max_records: int,
    max_expanded_bytes: int,
) -> tuple[str, int | None, bool]:
    try:
        reader = csv.DictReader(StringIO(text))
    except csv.Error as exc:
        raise DocumentParseError("CSV document is invalid") from exc
    if not reader.fieldnames:
        return text, None, False
    header = {str(key).strip().lower(): key for key in reader.fieldnames}
    time_key = _matching_key(header, _TIME_KEYS)
    sender_key = _matching_key(header, _SENDER_KEYS)
    content_key = _matching_key(header, _CONTENT_KEYS)
    if content_key is None or (time_key is None and sender_key is None):
        return text, None, False
    try:
        normalized = _stream_message_rows(
            reader,
            time_key=time_key,
            sender_key=sender_key,
            content_key=content_key,
            max_records=max_records,
            max_expanded_bytes=max_expanded_bytes,
            limit_label="CSV chat export",
        )
    except csv.Error as exc:
        raise DocumentParseError("CSV document is invalid") from exc
    if normalized is None:
        return text, None, False
    return normalized[0], normalized[1], True


def _stream_chat_json(
    text: str,
    *,
    max_records: int,
    max_expanded_bytes: int,
) -> tuple[str, int, bool] | None:
    decoder = json.JSONDecoder()
    index = _skip_ws(text, 0)
    if index >= len(text):
        return None
    if text[index] == "[":
        streamed = _stream_json_array(
            text,
            index,
            decoder,
            max_records=max_records,
            max_expanded_bytes=max_expanded_bytes,
        )
        return None if streamed is None else (streamed[0], streamed[1], True)
    if text[index] != "{":
        return None
    index += 1
    chat_result: tuple[str, int, bool] | None = None
    while True:
        index = _skip_ws(text, index)
        if index >= len(text):
            raise DocumentParseError("JSON document is invalid")
        if text[index] == "}":
            return chat_result
        try:
            key, index = decoder.raw_decode(text, index)
        except json.JSONDecodeError as exc:
            raise DocumentParseError("JSON document is invalid") from exc
        index = _skip_ws(text, index)
        if index >= len(text) or text[index] != ":":
            raise DocumentParseError("JSON document is invalid")
        index = _skip_ws(text, index + 1)
        if (
            chat_result is None
            and key in {"messages", "records", "items", "data"}
            and index < len(text)
            and text[index] == "["
        ):
            streamed = _stream_json_array(
                text,
                index,
                decoder,
                max_records=max_records,
                max_expanded_bytes=max_expanded_bytes,
            )
            if streamed is None:
                try:
                    _, index = decoder.raw_decode(text, index)
                except json.JSONDecodeError as exc:
                    raise DocumentParseError("JSON document is invalid") from exc
            else:
                chat_result = (streamed[0], streamed[1], True)
                index = streamed[2]
        else:
            try:
                _, index = decoder.raw_decode(text, index)
            except json.JSONDecodeError as exc:
                raise DocumentParseError("JSON document is invalid") from exc
        index = _skip_ws(text, index)
        if index < len(text) and text[index] == ",":
            index += 1
            continue
        if index < len(text) and text[index] == "}":
            return chat_result
        raise DocumentParseError("JSON document is invalid")


def _stream_json_array(
    text: str,
    start: int,
    decoder: json.JSONDecoder,
    *,
    max_records: int,
    max_expanded_bytes: int,
) -> tuple[str, int, int] | None:
    index = _skip_ws(text, start + 1)
    if index < len(text) and text[index] == "]":
        return None
    time_key: Any = None
    sender_key: Any = None
    content_key: Any = None
    output: list[str] = []
    expanded = 0
    count = 0
    while True:
        index = _skip_ws(text, index)
        if index >= len(text):
            raise DocumentParseError("JSON document is invalid")
        if text[index] == "]":
            break
        try:
            item, index = decoder.raw_decode(text, index)
        except json.JSONDecodeError as exc:
            raise DocumentParseError("JSON document is invalid") from exc
        if not isinstance(item, dict):
            return None
        if content_key is None:
            first = {str(key).strip().lower(): key for key in item}
            time_key = _matching_key(first, _TIME_KEYS)
            sender_key = _matching_key(first, _SENDER_KEYS)
            content_key = _matching_key(first, _CONTENT_KEYS)
            if content_key is None or (time_key is None and sender_key is None):
                return None
        count += 1
        if count > max_records:
            raise DocumentParseError("JSON chat export exceeds the configured record limit")
        line = _format_message_row(item, time_key, sender_key, content_key)
        expanded += len(line.encode("utf-8")) + (1 if output else 0)
        if expanded > max_expanded_bytes:
            raise DocumentParseError("extracted document text exceeds the configured limit")
        output.append(line)
        index = _skip_ws(text, index)
        if index < len(text) and text[index] == ",":
            index += 1
            continue
        if index < len(text) and text[index] == "]":
            break
        raise DocumentParseError("JSON document is invalid")
    if not output:
        return None
    return "\n".join(output), count, index + 1


def _stream_message_rows(
    rows: Any,
    *,
    time_key: Any,
    sender_key: Any,
    content_key: Any,
    max_records: int,
    max_expanded_bytes: int,
    limit_label: str,
) -> tuple[str, int] | None:
    output: list[str] = []
    expanded = 0
    count = 0
    for row in rows:
        if not isinstance(row, dict):
            return None
        count += 1
        if count > max_records:
            raise DocumentParseError(f"{limit_label} exceeds the configured record limit")
        line = _format_message_row(row, time_key, sender_key, content_key)
        expanded += len(line.encode("utf-8")) + (1 if output else 0)
        if expanded > max_expanded_bytes:
            raise DocumentParseError("extracted document text exceeds the configured limit")
        output.append(line)
    if not output:
        return None
    return "\n".join(output), count


def _format_message_row(
    row: dict[str, Any],
    time_key: Any,
    sender_key: Any,
    content_key: Any,
) -> str:
    timestamp = _format_timestamp(row.get(time_key)) if time_key is not None else ""
    sender = _sender_value(row.get(sender_key)) if sender_key is not None else "unknown"
    content = _content_value(row.get(content_key))
    prefix = f"[{timestamp}] " if timestamp else ""
    return f"{prefix}{sender or 'unknown'}: {content}"


def _skip_ws(text: str, index: int) -> int:
    length = len(text)
    while index < length and text[index] in " \t\r\n":
        index += 1
    return index


def _matching_key(index: dict[str, Any], aliases: Iterable[str]) -> Any | None:
    for alias in aliases:
        if alias in index:
            return index[alias]
    return None


def _format_timestamp(value: Any) -> str:
    if isinstance(value, bool):
        return ""
    if isinstance(value, (int, float)) and value >= 1_000_000_000:
        return datetime.fromtimestamp(int(value), tz=_SHANGHAI).strftime("%Y-%m-%d %H:%M:%S")
    return _string_value(value)


def _sender_value(value: Any) -> str:
    if isinstance(value, dict):
        for key in ("nickname", "card", "name", "nick", "user_id", "uin", "qq"):
            if value.get(key) not in (None, ""):
                return _string_value(value.get(key))
        return ""
    return _string_value(value)


def _content_value(value: Any) -> str:
    if isinstance(value, list):
        parts: list[str] = []
        for item in value:
            if isinstance(item, str):
                cleaned = item.strip()
                if cleaned:
                    parts.append(cleaned)
                continue
            if not isinstance(item, dict):
                continue
            kind = str(item.get("type") or "")
            data = item.get("data") if isinstance(item.get("data"), dict) else {}
            if kind == "text":
                text = _string_value(data.get("text") or item.get("text"))
                if text:
                    parts.append(text)
            elif kind:
                parts.append(f"[{kind}]")
        return " ".join(parts)
    return _string_value(value)


def _string_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.replace("\r\n", "\n").replace("\r", "\n").strip()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).strip()
